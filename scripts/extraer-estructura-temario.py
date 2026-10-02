#!/usr/bin/env python3
"""
Extrae la organización del temario del PowerPoint de organización a un JSON
que consume la app móvil:

    oposicion/organizacion/estructura_temario.ppsx  ->  estructura_temario.json

Qué se extrae:
  - Bloques del tercer ejercicio con su color (diapositiva "BLOQUES 3ER
    EJERCICIO") y el color de cada tema (diapositivas "Listado de temas").
  - Bloques del cuarto ejercicio (hoja "Listado de temas" del Excel de
    organización), con el color de su parte en el PowerPoint.
  - Los esquemas con la posición de cada tema, los marcos de bloque y las
    flechas que conectan temas y bloques (tercero, cuarto y combinado).
  - Las conexiones de CONEXIONES_EXTRA, que no están en el PowerPoint.
  - La idea clave de cada tema (texto que acompaña a su casilla en las
    diapositivas de detalle de cada bloque).

Uso (hay que repetirlo cuando cambie el PowerPoint):
    pip install python-pptx openpyxl
    python3 scripts/extraer-estructura-temario.py

Autor: Víctor Gutiérrez Marcos
"""

import io
import json
import re
import unicodedata
import warnings
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from pptx import Presentation

RAIZ = Path(__file__).resolve().parent.parent
PPSX = RAIZ / 'oposicion/organizacion/estructura_temario.ppsx'
EXCEL = RAIZ / 'oposicion/organizacion/preparacion_oposicion_tcee.xlsm'
TEMARIO = RAIZ / 'oposicion/temario/temario.json'
SALIDA = RAIZ / 'oposicion/organizacion/estructura_temario.json'

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
CODIGO = re.compile(r'^(\d)\.([ABC])\.(\d+)$')

# Diapositivas de la presentación (numeradas desde 1).
DIAP_LISTADO_3 = 4
DIAP_LISTADO_4 = 6
DIAP_BLOQUES_3 = 13
DIAPS_DETALLE = range(16, 40)
ESQUEMAS = [
    ('tercero', 'Tercer ejercicio', 5),
    ('cuarto', 'Cuarto ejercicio', 7),
    ('combinado', 'Tercer y cuarto ejercicio', 10),
]
# Conexiones que el autor quiere en la app y que no están dibujadas en el
# PowerPoint. Se añaden a todos los esquemas en los que aparecen los dos temas.
CONEXIONES_EXTRA = [
    {'de': '3.B.2', 'a': '3.B.3', 'flechaDe': True, 'flechaA': True},
]
CATEGORIAS_3 = ['MICROECONOMÍA', 'MACROECONOMÍA', 'MIXTO']
NOMBRES_CATEGORIA = {'MICROECONOMÍA': 'Microeconomía', 'MACROECONOMÍA': 'Macroeconomía', 'MIXTO': 'Mixto'}


def abrir_presentacion():
    """python-pptx no abre .ppsx: se cambia el tipo de contenido en memoria."""
    origen = zipfile.ZipFile(PPSX)
    tema = origen.read('ppt/theme/theme1.xml').decode('utf8')
    memoria = io.BytesIO()
    with zipfile.ZipFile(memoria, 'w', zipfile.ZIP_DEFLATED) as destino:
        for nombre in origen.namelist():
            datos = origen.read(nombre)
            if nombre == '[Content_Types].xml':
                datos = datos.replace(b'presentationml.slideshow.main+xml', b'presentationml.presentation.main+xml')
            destino.writestr(nombre, datos)
    colores_tema = dict(re.findall(r'<a:(accent\d|dk1|dk2|lt1|lt2)>.*?(?:srgbClr val|lastClr)="([0-9A-Fa-f]{6})"', tema, re.S))
    memoria.seek(0)
    return Presentation(memoria), colores_tema


def slug(texto):
    sin = unicodedata.normalize('NFD', texto.lower())
    sin = ''.join(c for c in sin if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^a-z0-9]+', '-', sin).strip('-')


def texto_de(forma):
    return ' '.join(forma.text_frame.text.split()) if forma.has_text_frame else ''


def codigo_de(texto):
    m = CODIGO.match(texto)
    return f'{m.group(1)}.{m.group(2)}.{int(m.group(3))}' if m else None


class Lector:
    def __init__(self, prs, colores_tema):
        self.prs = prs
        self.colores_tema = colores_tema
        self.ancho = prs.slide_width
        self.alto = prs.slide_height

    def color(self, elemento_relleno):
        """Color de un <a:solidFill> (o equivalente) como RRGGBB."""
        if elemento_relleno is None:
            return None
        rgb = elemento_relleno.find('a:srgbClr', NS)
        if rgb is not None:
            return rgb.get('val').upper()
        esquema = elemento_relleno.find('a:schemeClr', NS)
        if esquema is not None:
            return self.colores_tema.get(esquema.get('val'), '').upper() or None
        return None

    def relleno(self, forma):
        """(color, tramado): tramado indica una casilla repetida o de otro ejercicio."""
        sppr = forma._element.find('p:spPr', NS)
        if sppr is None:
            return None, False
        trama = sppr.find('a:pattFill', NS)
        if trama is not None:
            return self.color(trama.find('a:fgClr', NS)), True
        solido = sppr.find('a:solidFill', NS)
        if solido is not None:
            return self.color(solido), solido.find('.//a:alpha', NS) is not None
        return None, False

    def formas(self, diapositiva):
        """Todas las formas con su rectángulo en coordenadas de la diapositiva
        (deshaciendo la transformación de los grupos)."""
        salida = []

        def recorrer(formas, transformar):
            for f in formas:
                if type(f).__name__ == 'GroupShape':
                    x = f._element.find('p:grpSpPr/a:xfrm', NS)
                    off, ext = x.find('a:off', NS), x.find('a:ext', NS)
                    choff, chext = x.find('a:chOff', NS), x.find('a:chExt', NS)
                    ox, oy, ex, ey = int(off.get('x')), int(off.get('y')), int(ext.get('cx')), int(ext.get('cy'))
                    cx, cy, cw, ch = int(choff.get('x')), int(choff.get('y')), int(chext.get('cx')), int(chext.get('cy'))

                    def hijo(r, t=transformar, ox=ox, oy=oy, ex=ex, ey=ey, cx=cx, cy=cy, cw=cw, ch=ch):
                        fx, fy = (ex / cw if cw else 1), (ey / ch if ch else 1)
                        return t((ox + (r[0] - cx) * fx, oy + (r[1] - cy) * fy, r[2] * fx, r[3] * fy))

                    recorrer(f.shapes, hijo)
                elif f.left is not None and f.width is not None:
                    salida.append((f, transformar((f.left, f.top, f.width, f.height))))

        recorrer(diapositiva.shapes, lambda r: r)
        return salida

    def norm(self, r):
        return [round(r[0] / self.ancho, 4), round(r[1] / self.alto, 4), round(r[2] / self.ancho, 4), round(r[3] / self.alto, 4)]


def distancia(a, b):
    dx = max(a[0] - (b[0] + b[2]), b[0] - (a[0] + a[2]), 0)
    dy = max(a[1] - (b[1] + b[3]), b[1] - (a[1] + a[3]), 0)
    return (dx * dx + dy * dy) ** 0.5


def extraer_bloques(lector):
    prs = lector.prs
    # Color de cada tema en los listados.
    color_tema = {}
    for n in (DIAP_LISTADO_3, DIAP_LISTADO_4):
        for f, _ in lector.formas(prs.slides[n - 1]):
            c = codigo_de(texto_de(f))
            if c:
                color_tema[c] = lector.relleno(f)[0]

    # Bloques del tercer ejercicio: nombre, color y categoría (la cabecera de su columna).
    cabeceras, cajas = [], []
    for f, r in lector.formas(prs.slides[DIAP_BLOQUES_3 - 1]):
        t = texto_de(f)
        if t in CATEGORIAS_3:
            cabeceras.append((t, r))
        elif t and not t.upper().startswith('BLOQUES') and t != 'Economía':
            color = lector.relleno(f)[0]
            if color:
                cajas.append((t, color, r))
    bloques = []
    for nombre, color, r in sorted(cajas, key=lambda c: (min(cabeceras, key=lambda k: abs((k[1][0] + k[1][2] / 2) - (c[2][0] + c[2][2] / 2)))[1][0], c[2][1])):
        categoria = min(cabeceras, key=lambda k: abs((k[1][0] + k[1][2] / 2) - (r[0] + r[2] / 2)))[0]
        temas = sorted((t for t, c in color_tema.items() if c == color and t.startswith('3.')), key=orden_tema)
        bloques.append({'id': slug(nombre), 'nombre': nombre, 'ejercicio': 3, 'categoria': NOMBRES_CATEGORIA[categoria], 'color': color, 'temas': temas})

    # Bloques del cuarto ejercicio: los del Excel, con el color de su parte en el PowerPoint.
    warnings.simplefilter('ignore')
    hoja = openpyxl.load_workbook(EXCEL)['Listado de temas']
    del_cuarto = defaultdict(list)
    for fila in range(3, hoja.max_row + 1):
        if hoja[f'B{fila}'].value == 4:
            clave = (hoja[f'C{fila}'].value, hoja[f'G{fila}'].value, hoja[f'H{fila}'].value)
            del_cuarto[clave].append(f"4.{hoja[f'C{fila}'].value}.{hoja[f'D{fila}'].value}")
    for (parte, categoria, nombre), temas in del_cuarto.items():
        bloques.append({
            'id': slug(f'4{parte}-{nombre}'), 'nombre': nombre, 'ejercicio': 4,
            'categoria': categoria, 'color': color_tema[temas[0]], 'temas': sorted(temas, key=orden_tema),
        })
    return bloques, color_tema


def orden_tema(codigo):
    e, p, n = codigo.split('.')
    return (int(e), p, int(n))


def extraer_esquema(lector, numero, bloques, color_tema):
    por_nombre = {b['nombre']: b['id'] for b in bloques}
    por_nombre['Historia del pensamiento económico'] = por_nombre.get('Historia del pensamiento económico', por_nombre.get('Pensamiento económico'))
    por_color = {b['color']: b['id'] for b in bloques if b['ejercicio'] == 3}
    color_bloque = {b['id']: b['color'] for b in bloques}
    diapositiva = lector.prs.slides[numero - 1]

    nodos, marcos, etiquetas = [], [], []
    por_id = {}  # id de forma -> ('tema', código) | ('bloque', id)
    rects = []   # (rectángulo, destino) para resolver extremos sueltos
    conectores = []
    for f, r in lector.formas(diapositiva):
        tipo = type(f).__name__
        if tipo == 'Connector':
            conectores.append((f, r))
            continue
        t = texto_de(f)
        c = codigo_de(t)
        if c:
            color, tramado = lector.relleno(f)
            nodos.append({'tema': c, 'r': lector.norm(r), **({'repetido': True} if tramado else {})})
            por_id[str(f.shape_id)] = ('tema', c)
            rects.append((r, ('tema', c)))
        elif f.name.startswith('!!') and t in por_nombre:
            etiquetas.append({'bloque': por_nombre[t], 'r': lector.norm(r)})
            por_id[str(f.shape_id)] = ('bloque', por_nombre[t])
            rects.append((r, ('bloque', por_nombre[t])))
        elif 'esquinas redondeadas' in f.name or f.name.startswith('Rectángulo'):
            ln = f._element.find('p:spPr/a:ln', NS)
            color = lector.color(ln.find('a:solidFill', NS)) if ln is not None else None
            if color:
                marcos.append((f, r, color))

    # Cada marco es el de un bloque: el de su color de línea o, si no, el de la mayoría de los temas que contiene.
    marcos_json = []
    for f, r, color in marcos:
        dentro = Counter()
        for n in nodos:
            x, y = n['r'][0] * lector.ancho, n['r'][1] * lector.alto
            if r[0] <= x <= r[0] + r[2] and r[1] <= y <= r[1] + r[3] and not n.get('repetido'):
                dentro[color_tema.get(n['tema'])] += 1
        bloque = por_color.get(color) or (por_color.get(dentro.most_common(1)[0][0]) if dentro else None)
        if bloque:
            por_id[str(f.shape_id)] = ('bloque', bloque)
            # Un bloque puede tener marcos interiores: se dibuja solo el mayor.
            previo = next((m for m in marcos_json if m['bloque'] == bloque), None)
            if previo is None:
                marcos_json.append({'bloque': bloque, 'r': lector.norm(r)})
            elif r[2] * r[3] > previo['r'][2] * previo['r'][3] * lector.ancho * lector.alto:
                previo['r'] = lector.norm(r)

    def extremo(conector, r, cual):
        e = conector._element.find(f'.//a:{cual}', NS)
        if e is not None and e.get('id') in por_id:
            return por_id[e.get('id')]
        # Extremo sin enganchar: la casilla más cercana a ese punto del conector.
        xfrm = conector._element.find('p:spPr/a:xfrm', NS)
        vol_h, vol_v = xfrm.get('flipH') == '1', xfrm.get('flipV') == '1'
        inicio = cual == 'stCxn'
        x = r[0] + (r[2] if (inicio == vol_h) else 0)
        y = r[1] + (r[3] if (inicio == vol_v) else 0)
        mejor = min(rects, key=lambda q: distancia((x, y, 0, 0), q[0]), default=None)
        if mejor and distancia((x, y, 0, 0), mejor[0]) < lector.ancho * 0.012:
            return mejor[1]
        return None

    conexiones, vistas = [], set()
    for f, r in conectores:
        de, a = extremo(f, r, 'stCxn'), extremo(f, r, 'endCxn')
        if not de or not a or de == a:
            continue
        ln = f._element.find('p:spPr/a:ln', NS)
        cabeza = ln.find('a:headEnd', NS) if ln is not None else None
        cola = ln.find('a:tailEnd', NS) if ln is not None else None
        trazo = ln.find('a:prstDash', NS) if ln is not None else None
        clave = (de, a)
        if clave in vistas or (a, de) in vistas:
            continue
        vistas.add(clave)
        conexiones.append({
            'de': {de[0]: de[1]}, 'a': {a[0]: a[1]},
            'flechaDe': cabeza is not None and cabeza.get('type') not in (None, 'none'),
            'flechaA': cola is not None and cola.get('type') not in (None, 'none'),
            'color': (lector.color(ln.find('a:solidFill', NS)) if ln is not None else None) or (color_tema.get(de[1]) if de[0] == 'tema' else color_bloque.get(de[1])),
            'discontinua': trazo is not None and trazo.get('val') != 'solid',
        })
    presentes = {n['tema'] for n in nodos}
    for extra in CONEXIONES_EXTRA:
        de, a = ('tema', extra['de']), ('tema', extra['a'])
        if extra['de'] in presentes and extra['a'] in presentes and (de, a) not in vistas and (a, de) not in vistas:
            vistas.add((de, a))
            conexiones.append({
                'de': {'tema': extra['de']}, 'a': {'tema': extra['a']},
                'flechaDe': extra['flechaDe'], 'flechaA': extra['flechaA'],
                'color': color_tema.get(extra['de']), 'discontinua': False,
            })
    return {'nodos': nodos, 'marcos': marcos_json, 'etiquetas': etiquetas, 'conexiones': conexiones}


def palabras(texto):
    """Raíces (6 letras) de las palabras significativas de un texto, para comparar."""
    sin = unicodedata.normalize('NFD', texto.lower())
    sin = ''.join(c for c in sin if unicodedata.category(c) != 'Mn')
    return {p[:6] for p in re.findall(r'[a-zñ]+', sin) if len(p) >= 5}


def extraer_ideas(lector):
    """Idea clave de cada tema: el texto breve que lo acompaña en las diapositivas de detalle.

    Los textos no están enganchados a las casillas, así que se emparejan por
    cercanía y solo se aceptan si comparten vocabulario con el título oficial
    del tema (temario.json); así no se atribuye a un tema el rótulo de otro."""
    temario = json.loads(TEMARIO.read_text(encoding='utf8'))
    titulos = {t['codigo']: t['titulo'] for e in temario['ejercicios'] for p in e['partes'] for t in p['temas']}
    ideas = {}
    for n in DIAPS_DETALLE:
        temas, textos = [], []
        for f, r in lector.formas(lector.prs.slides[n - 1]):
            if type(f).__name__ == 'Connector' or not f.has_text_frame:
                continue
            t = texto_de(f)
            if not t:
                continue
            c = codigo_de(t)
            if c:
                if not lector.relleno(f)[1]:
                    temas.append((c, r))
            elif 'Título' not in f.name and len(t) > 3 and not CODIGO.match(t.split()[0]):
                textos.append((t, r))
        for t, q in textos:
            mejor = None
            for c, r in temas:
                d = distancia(r, q)
                comunes = len(palabras(t) & palabras(titulos.get(c, '')))
                if d < lector.ancho * 0.03 and comunes and (mejor is None or (comunes, -d) > (mejor[0], -mejor[1])):
                    mejor = (comunes, d, c)
            if mejor:
                comunes, d, c = mejor
                if c not in ideas or (comunes, -d) > (ideas[c][0], -ideas[c][1]):
                    ideas[c] = (comunes, d, t.strip(' .') + ('.' if len(t) > 40 else ''))

    def aporta(c, idea):
        # Una idea que repite casi entero el título oficial no añade nada.
        titulo = titulos[c]
        return not (len(idea) >= 0.6 * len(titulo) and len(palabras(idea) & palabras(titulo)) >= 0.8 * len(palabras(idea)))

    return {c: v[2] for c, v in sorted(ideas.items(), key=lambda kv: orden_tema(kv[0])) if aporta(c, v[2])}


def main():
    prs, colores_tema = abrir_presentacion()
    lector = Lector(prs, colores_tema)
    bloques, color_tema = extraer_bloques(lector)
    esquemas = []
    for ident, titulo, numero in ESQUEMAS:
        esquemas.append({'id': ident, 'titulo': titulo, **extraer_esquema(lector, numero, bloques, color_tema)})
    datos = {
        '_comentario': 'Organización del temario extraída de estructura_temario.ppsx con scripts/extraer-estructura-temario.py. No editar a mano. Autor: Víctor Gutiérrez Marcos.',
        'version': 1,
        'proporcion': round(lector.ancho / lector.alto, 4),
        'bloques': bloques,
        'ideas': extraer_ideas(lector),
        'esquemas': esquemas,
    }
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
    print(f'✔ {SALIDA.relative_to(RAIZ)}')
    print(f"  {len(bloques)} bloques, {sum(len(b['temas']) for b in bloques)} temas, {len(datos['ideas'])} ideas clave")
    for e in esquemas:
        print(f"  {e['id']}: {len(e['nodos'])} casillas, {len(e['marcos'])} marcos, {len(e['conexiones'])} conexiones")


if __name__ == '__main__':
    main()
