#!/usr/bin/env python3
"""
Extrae la organización del temario de DCE de la presentación de Manuel Cabado
García al mismo JSON que la de TCEE (scripts/extraer-estructura-temario.py),
para que la app muestre bloques con colores, esquemas y conexiones, y ordene
los cronogramas siguiéndolos:

    app/para-manuel/fuentes/organizacion_temario_DCE.pptx
        ->  app/para-manuel/oposicion/organizacion/estructura_temario.json

Manuel copia la carpeta app/para-manuel/oposicion/ a su web.

Qué se extrae:
  - Bloques del 1.er y del 3.er ejercicio con su color y su categoría
    (diapositivas «Bloques»).
  - Los esquemas «Esquema por temas» con la posición de cada tema, los marcos
    de bloque y las flechas entre temas y bloques.
  - Las conexiones «CONECTADO CON» de la ficha de cada tema que no estén
    dibujadas en el esquema, en «conexiones» (la app no las dibuja).
  - La idea clave de cada tema: el título corto de su ficha (solo en el 3.º).

Uso (hay que repetirlo cuando cambie la presentación):
    pip install python-pptx
    python3 scripts/extraer-estructura-dce.py

Autor: Víctor Gutiérrez Marcos
"""

import importlib.util
import json
import re
from collections import Counter
from pathlib import Path

from pptx import Presentation

RAIZ = Path(__file__).resolve().parent.parent
PPTX = RAIZ / 'app/para-manuel/fuentes/organizacion_temario_DCE.pptx'
TEMARIO = RAIZ / 'app/para-manuel/oposicion/temario/temario.json'
SALIDA = RAIZ / 'app/para-manuel/oposicion/organizacion/estructura_temario.json'

# Se reutiliza la lectura de formas y colores del extractor de TCEE.
_spec = importlib.util.spec_from_file_location('estructura_tcee', RAIZ / 'scripts/extraer-estructura-temario.py')
tcee = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(tcee)
NS, Lector, texto_de, codigo_de, slug, distancia, orden_tema = (
    tcee.NS, tcee.Lector, tcee.texto_de, tcee.codigo_de, tcee.slug, tcee.distancia, tcee.orden_tema)

# Diapositivas de la presentación (numeradas desde 1).
EJERCICIOS = {
    # ejercicio: (esquema, bloques, fichas, id del esquema, título)
    1: (7, 8, range(53, 89), 'primero', 'Primer ejercicio'),
    3: (4, 5, range(9, 53), 'tercero', 'Tercer ejercicio'),
}
CATEGORIAS = {
    'MICROECONOMÍA': 'Microeconomía', 'MACROECONOMÍA': 'Macroeconomía',
    'ECONOMÍA INTERNACIONAL': 'Economía internacional', 'MIXTO': 'Mixto',
    'ECONOMÍA ESPAÑOLA': 'Economía española', 'ECONOMÍA PÚBLICA': 'Economía pública',
    'POLÍTICAS COMUNITARIAS': 'Políticas comunitarias', 'INSTITUCIONES MULTILATERALES': 'Instituciones multilaterales',
}
BOTONES = {'Menú', 'Esquema', 'Listado', 'Bloques'}
LISTA_TEMAS = re.compile(r'^[ABC]\d+( · [ABC]\d+)*$')


def extraer_bloques(lector, ejercicio, numero):
    """Recorre la diapositiva en orden: rótulo de categoría y caja de cada bloque
    (con su color), con el nombre en el primer párrafo y la lista «A3 · A4 · …»
    de sus temas en el segundo."""
    bloques, categoria = [], None
    for f, _ in lector.formas(lector.prs.slides[numero - 1]):
        t = texto_de(f)
        if t in CATEGORIAS:
            categoria = CATEGORIAS[t]
            continue
        if not categoria or not t or t in BOTONES:
            continue
        parrafos = [' '.join(p.text.split()) for p in f.text_frame.paragraphs if p.text.strip()]
        color = lector.relleno(f)[0]
        if len(parrafos) == 2 and LISTA_TEMAS.match(parrafos[1]) and color:
            nombre, lista = parrafos
            bloques.append({
                'id': f'e{ejercicio}-{slug(nombre)}', 'nombre': nombre, 'ejercicio': ejercicio, 'categoria': categoria, 'color': color,
                'temas': sorted((f'{ejercicio}.{x[0]}.{int(x[1:])}' for x in lista.split(' · ')), key=orden_tema),
            })
    return bloques


def extraer_esquema(lector, numero, bloques):
    """Casillas de tema, marcos de bloque (rectángulos tramados), rótulos y las
    líneas que los unen (formas «line», sin enganchar: se resuelven por cercanía)."""
    de_tema = {t: b for b in bloques for t in b['temas']}
    por_nombre = {b['nombre']: b for b in bloques}
    nodos, marcos, etiquetas, lineas = [], [], [], []
    rect_nodo, rect_bloque = [], []
    for f, r in lector.formas(lector.prs.slides[numero - 1]):
        t = texto_de(f)
        c = codigo_de(t)
        geom = f._element.find('p:spPr/a:prstGeom', NS)
        if c:
            nodos.append({'tema': c, 'r': lector.norm(r)})
            rect_nodo.append((r, ('tema', c)))
        elif geom is not None and geom.get('prst') == 'line':
            lineas.append((f, r))
        elif t and t.startswith(tuple(por_nombre)):
            # El rótulo puede venir cortado: se busca el bloque por el comienzo.
            b = next(b for n, b in por_nombre.items() if n.startswith(t) or t.startswith(n))
            etiquetas.append({'bloque': b['id'], 'r': lector.norm(r)})
            rect_bloque.append((r, ('bloque', b['id'])))
        elif not t and lector.relleno(f)[1]:
            marcos.append(r)

    # Cada marco es el del bloque de la mayoría de los temas que contiene.
    marcos_json = []
    for r in marcos:
        dentro = Counter()
        for n in nodos:
            x, y = n['r'][0] * lector.ancho, n['r'][1] * lector.alto
            if r[0] <= x <= r[0] + r[2] and r[1] <= y <= r[1] + r[3]:
                dentro[de_tema[n['tema']]['id']] += 1
        if dentro:
            bloque = dentro.most_common(1)[0][0]
            marcos_json.append({'bloque': bloque, 'r': lector.norm(r)})
            rect_bloque.append((r, ('bloque', bloque)))

    def extremo(f, r, inicio):
        xfrm = f._element.find('p:spPr/a:xfrm', NS)
        vol_h, vol_v = xfrm.get('flipH') == '1', xfrm.get('flipV') == '1'
        x = r[0] + (r[2] if (inicio == vol_h) else 0)
        y = r[1] + (r[3] if (inicio == vol_v) else 0)
        # Antes una casilla de tema (están dentro de los marcos) que un bloque.
        for candidatos, tope in ((rect_nodo, 0.006), (rect_bloque, 0.015)):
            mejor = min(candidatos, key=lambda q: distancia((x, y, 0, 0), q[0]), default=None)
            if mejor and distancia((x, y, 0, 0), mejor[0]) <= lector.ancho * tope:
                return mejor[1]
        return None

    color_de = {('tema', t): b['color'] for t, b in de_tema.items()} | {('bloque', b['id']): b['color'] for b in bloques}
    conexiones, vistas, sueltas = [], set(), 0
    for f, r in lineas:
        de, a = extremo(f, r, True), extremo(f, r, False)
        if not de or not a or de == a:
            sueltas += 1
            continue
        if (de, a) in vistas or (a, de) in vistas:
            continue
        vistas.add((de, a))
        ln = f._element.find('p:spPr/a:ln', NS)
        cabeza, cola = ln.find('a:headEnd', NS), ln.find('a:tailEnd', NS)
        trazo = ln.find('a:prstDash', NS)
        conexiones.append({
            'de': {de[0]: de[1]}, 'a': {a[0]: a[1]},
            'flechaDe': cabeza is not None and cabeza.get('type') not in (None, 'none'),
            'flechaA': cola is not None and cola.get('type') not in (None, 'none'),
            'color': lector.color(ln.find('a:solidFill', NS)) or color_de[de],
            'discontinua': trazo is not None and trazo.get('val') != 'solid',
        })
    return {'nodos': nodos, 'marcos': marcos_json, 'etiquetas': etiquetas, 'conexiones': conexiones}, sueltas


def extraer_fichas(lector, fichas, bloques):
    """Título corto (idea) y «CONECTADO CON» de la ficha de cada tema."""
    por_nombre = {b['nombre']: b['id'] for b in bloques}
    ideas, conectados = {}, {}
    for n in fichas:
        textos = [texto_de(f) for f, _ in lector.formas(lector.prs.slides[n - 1])]
        textos = [t for t in textos if t]
        i = next(i for i, t in enumerate(textos) if codigo_de(t))
        tema = codigo_de(textos[i])
        if textos[i + 2] != 'BLOQUE':
            ideas[tema] = textos[i + 1]
        destino = []
        for t in textos[textos.index('CONECTADO CON') + 1:]:
            if t in BOTONES or t.startswith('Leer los apuntes'):
                break
            if codigo_de(t):
                destino.append(('tema', codigo_de(t)))
            elif t in por_nombre:
                destino.append(('bloque', por_nombre[t]))
            else:
                raise SystemExit(f'✘ {tema}: «{t}» en CONECTADO CON no es un tema ni un bloque')
        conectados[tema] = destino
    return ideas, conectados


def main():
    prs = Presentation(PPTX)
    lector = Lector(prs, {})
    bloques, ideas, esquemas, sin_dibujar, avisos = [], {}, [], [], []
    for ejercicio, (n_esquema, n_bloques, fichas, ident, titulo) in EJERCICIOS.items():
        del_ej = extraer_bloques(lector, ejercicio, n_bloques)
        bloques += del_ej
        esquema, sueltas = extraer_esquema(lector, n_esquema, del_ej)
        if sueltas:
            avisos.append(f'{ident}: {sueltas} líneas sin enganchar a un tema o bloque')
        idea, conectados = extraer_fichas(lector, fichas, del_ej)
        ideas |= idea
        # Lo de «CONECTADO CON» que no está dibujado va aparte, sin dibujar
        # (llenaría el esquema de líneas), para «Relacionados» y el cronograma.
        color_de = {('tema', t): b['color'] for b in del_ej for t in b['temas']} | {('bloque', b['id']): b['color'] for b in del_ej}
        dibujadas = {(tuple(c['de'].items())[0], tuple(c['a'].items())[0]) for c in esquema['conexiones']}
        for tema, destinos in conectados.items():
            for d in destinos:
                de = ('tema', tema)
                if de == d or (de, d) in dibujadas or (d, de) in dibujadas or d not in color_de:
                    continue
                dibujadas.add((de, d))
                sin_dibujar.append({'de': {'tema': tema}, 'a': {d[0]: d[1]}, 'flechaDe': False, 'flechaA': False, 'color': color_de[d], 'discontinua': True})
        esquemas.append({'id': ident, 'titulo': titulo, **esquema})

    # Comprobaciones: cada tema del temario está en un solo bloque y en el esquema.
    temario = json.loads(TEMARIO.read_text(encoding='utf8'))
    codigos = {t['codigo'] for e in temario['ejercicios'] if e['id'] in EJERCICIOS for p in e.get('partes', []) for t in p['temas']}
    en_bloques = Counter(t for b in bloques for t in b['temas'])
    errores = [f'{t} está en {n} bloques' for t, n in en_bloques.items() if n != 1]
    errores += [f'{t} no está en ningún bloque' for t in sorted(codigos - set(en_bloques), key=orden_tema)]
    errores += [f'{t} no existe en el temario' for t in sorted(set(en_bloques) - codigos, key=orden_tema)]
    for e in esquemas:
        en_esquema = {n['tema'] for n in e['nodos']}
        errores += [f"{t} no está en el esquema {e['id']}" for t in sorted(codigos - en_esquema, key=orden_tema) if t.startswith(f"{next(k for k, v in EJERCICIOS.items() if v[3] == e['id'])}.")]
    if errores:
        raise SystemExit('✘ ' + '\n✘ '.join(errores))

    datos = {
        '_comentario': 'Organización del temario de DCE extraída de organizacion_temario_DCE.pptx (Manuel Cabado García) con scripts/extraer-estructura-dce.py. No editar a mano. Autor: Víctor Gutiérrez Marcos.',
        'version': 1,
        'proporcion': round(lector.ancho / lector.alto, 4),
        'bloques': bloques,
        'ideas': dict(sorted(ideas.items(), key=lambda kv: orden_tema(kv[0]))),
        'esquemas': esquemas,
        'conexiones': sin_dibujar,
    }
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    SALIDA.write_text(json.dumps(datos, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
    print(f'✔ {SALIDA.relative_to(RAIZ)}')
    print(f"  {len(bloques)} bloques, {sum(len(b['temas']) for b in bloques)} temas, {len(ideas)} ideas clave")
    for e in esquemas:
        print(f"  {e['id']}: {len(e['nodos'])} casillas, {len(e['marcos'])} marcos, {len(e['conexiones'])} conexiones")
    print(f'  {len(sin_dibujar)} conexiones de las fichas sin dibujar')
    for a in avisos:
        print(f'  ⚠ {a}')


if __name__ == '__main__':
    main()
