#!/usr/bin/env python3
"""
Crea oposicion/temario/latex/reference.docx, la plantilla de estilos que usa
pandoc para generar el Word de cada tema (construir-tema.py).

Parte de la plantilla por defecto de pandoc (que trae los nombres de estilo que
pandoc espera) y le aplica la estética de los temas en Word
(oposicion/organizacion/1_plantilla_temas_largos.dotx) y de la web: Palatino
Linotype, títulos en banda morada, cajas de nota y pies de imagen, cabecera y
pie con número de página.

    python3 scripts/temario/crear-reference-docx.py

Solo hace falta volver a ejecutarlo si cambia la estética.

Autor: Víctor Gutiérrez Marcos
"""
import os
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from comun import RAIZ, PANDOC, ruta_windows  # noqa: E402

SALIDA = os.path.join(RAIZ, 'oposicion', 'temario', 'latex', 'reference.docx')
MORADO = RGBColor(0x5F, 0x29, 0x87)
GRIS = RGBColor(0x76, 0x71, 0x71)
FUENTE = 'Palatino Linotype'


def sombreado(pPr, color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    pPr.append(shd)


def bordes(pPr, color, lados=('top', 'left', 'bottom', 'right'), grosor='4', espacio='4'):
    pbdr = OxmlElement('w:pBdr')
    for lado in lados:
        b = OxmlElement(f'w:{lado}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), grosor)
        b.set(qn('w:space'), espacio)
        b.set(qn('w:color'), color)
        pbdr.append(b)
    pPr.append(pbdr)


def fuente(estilo, tam=None, color=None, negrita=None, cursiva=None, mayus=None):
    f = estilo.font
    f.name = FUENTE
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for atr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(atr), FUENTE)
    for atr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(atr)) is not None:
            del rfonts.attrib[qn(atr)]
    if tam:
        f.size = Pt(tam)
    if color is not None:
        f.color.rgb = color
    if negrita is not None:
        f.bold = negrita
    if cursiva is not None:
        f.italic = cursiva
    if mayus is not None:
        f.all_caps = mayus


def estilo_parrafo(doc, nombre, base='Normal'):
    try:
        return doc.styles[nombre]
    except KeyError:
        e = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        e.base_style = doc.styles[base]
        e.quick_style = True
        return e


def campo(parrafo, instruccion):
    run = parrafo.add_run()
    for tipo, texto in (('begin', None), (None, instruccion), ('separate', None), (None, '1'), ('end', None)):
        if tipo:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), tipo)
            run._r.append(fc)
        elif texto == instruccion:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = f' {instruccion} '
            run._r.append(it)
        else:
            t = OxmlElement('w:t')
            t.text = texto
            run._r.append(t)
    return run


def main():
    with tempfile.TemporaryDirectory(dir=os.path.join(RAIZ, 'oposicion', 'temario', 'latex')) as tmp:
        base = os.path.join(tmp, 'base.docx')
        subprocess.run([PANDOC, '-o', ruta_windows(base), '--print-default-data-file', 'reference.docx'],
                       check=True)
        doc = Document(base)

    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.6), Cm(2.4)

    estilos = doc.styles
    for nombre in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
        e = estilos[nombre]
        fuente(e, tam=11)
        e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        e.paragraph_format.space_after = Pt(6)
    estilos['Compact'].paragraph_format.space_after = Pt(2)

    # Títulos
    t = estilos['Title']
    fuente(t, tam=13, color=MORADO, negrita=True, mayus=True)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h1 = estilos['Heading 1']
    fuente(h1, tam=12, color=RGBColor(0xFF, 0xFF, 0xFF), negrita=True, mayus=True)
    sombreado(h1.element.get_or_add_pPr(), '5F2987')
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    h2 = estilos['Heading 2']
    fuente(h2, tam=11, color=MORADO, negrita=True, mayus=False)
    sombreado(h2.element.get_or_add_pPr(), 'E2EFD9')
    h3 = estilos['Heading 3']
    fuente(h3, tam=11, color=MORADO, negrita=False, cursiva=True)
    for nombre in ('Heading 4', 'Heading 5', 'Heading 6'):
        try:
            fuente(estilos[nombre], tam=11, color=RGBColor(0x4A, 0x1F, 0x6B), negrita=True)
        except KeyError:
            pass

    for nombre in ('Author', 'Date', 'Subtitle'):
        try:
            e = estilos[nombre]
            fuente(e, tam=9, color=MORADO, negrita=False, cursiva=True, mayus=False)
            e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        except KeyError:
            pass

    # Figuras, notas y bibliografía
    for nombre in ('Figure', 'Captioned Figure'):
        try:
            estilos[nombre].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            estilos[nombre].paragraph_format.keep_with_next = True  # con su «Fuente»
        except KeyError:
            pass
    for nombre in ('Image Caption', 'Table Caption', 'Caption'):
        try:
            e = estilos[nombre]
            fuente(e, tam=10, color=MORADO, negrita=False, cursiva=False)
            e.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # El título va encima del gráfico: que no se quede solo al pie de página
            e.paragraph_format.keep_with_next = True
        except KeyError:
            pass
    fuente(estilos['Footnote Text'], tam=9)
    try:
        fuente(estilos['Bibliography'], tam=10)
    except KeyError:
        pass

    # Estilos propios (los nombres coinciden con los de la plantilla .dotx)
    nota = estilo_parrafo(doc, 'Nota al opositor')
    fuente(nota, tam=9.5)
    ppr = nota.element.get_or_add_pPr()
    bordes(ppr, '5F2987', lados=('left',), grosor='24', espacio='8')
    sombreado(ppr, 'F3EEF7')
    anot = estilo_parrafo(doc, 'Caja de anotaciones')
    fuente(anot, tam=9)
    ppr = anot.element.get_or_add_pPr()
    bordes(ppr, '767171')
    sombreado(ppr, 'E7E6E6')
    idea = estilo_parrafo(doc, 'Idea clave')
    fuente(idea, tam=11, cursiva=True)
    ppr = idea.element.get_or_add_pPr()
    bordes(ppr, 'C8D8C0')
    sombreado(ppr, 'E2EFD9')
    esq = estilo_parrafo(doc, 'Esquema')
    fuente(esq, tam=10)
    ppr = esq.element.get_or_add_pPr()
    bordes(ppr, '5F2987')
    sombreado(ppr, 'FAF9F6')
    pie = estilo_parrafo(doc, 'Pie de imagen')
    fuente(pie, tam=8, cursiva=True, color=RGBColor(0x3B, 0x38, 0x38))
    pie.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    try:
        ref = estilos['Referencia a tema']
    except KeyError:
        ref = estilos.add_style('Referencia a tema', WD_STYLE_TYPE.CHARACTER)
    ref.font.color.rgb = GRIS

    # Cabecera y pie
    cab = sec.header.paragraphs[0] if sec.header.paragraphs else sec.header.add_paragraph()
    cab.text = ''
    r = cab.add_run('Oposición a Técnico Comercial y Economista del Estado · Víctor Gutiérrez Marcos')
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = MORADO
    cab.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pie_p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
    pie_p.text = ''
    pie_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for parte in (campo(pie_p, 'PAGE'), pie_p.add_run('/'), campo(pie_p, 'NUMPAGES')):
        parte.font.size = Pt(9)
        parte.font.bold = True
        parte.font.color.rgb = MORADO

    doc.core_properties.author = 'Víctor Gutiérrez Marcos'
    doc.save(SALIDA)
    print('Escrito', SALIDA)


if __name__ == '__main__':
    main()
